"""Parse every downloaded attachment for a province and write the 16-column workbook."""
import sys, json, pathlib, collections, datetime
sys.path.insert(0, str(pathlib.Path(__file__).parent))
import parse_xml, parse_xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

COLS = ['Tỉnh/TP hiện hành','Tỉnh/TP theo nguồn','Số QĐ/Văn bản','Ngày tài liệu','Kỳ dữ liệu',
        'Cơ quan','Phạm vi','Nội dung/Bảng','Chỉ tiêu','Loại số liệu','Giá trị gốc',
        'Giá trị chuẩn hóa','ĐVT','Quy đổi VND','Nguồn','Ghi chú']
KEY = ['Tỉnh/TP theo nguồn','Số QĐ/Văn bản','Kỳ dữ liệu','Nội dung/Bảng','Chỉ tiêu','Loại số liệu']

def build(province, workdir):
    out = pathlib.Path(workdir)
    cat = {c['ID']: c for c in json.load(open(out / 'catalog.json'))}
    files = json.load(open(out / 'files.json'))

    by_report = collections.defaultdict(list)
    for f in files: by_report[f['report_id']].append(f)

    rows, status = [], {}
    for rid, fs in by_report.items():
        rec = cat.get(rid)
        if not rec: continue
        xml = [f for f in fs if f['filename'].lower().endswith('.xml')]
        sheets = [f for f in fs if f['filename'].lower().endswith(('.xls', '.xlsx'))]
        got, how, err = 0, '', ''
        for f in xml:
            try:
                r = list(parse_xml.rows_from(out / f['path'], rec, province))
                rows += r; got += len(r); how = 'XML'
            except Exception as e: err = f"{type(e).__name__}: {e}"
        if not got:
            for f in sheets:
                try:
                    r = list(parse_xlsx.rows_from(out / f['path'], rec, province))
                    rows += r; got += len(r); how = 'XLSX'
                except Exception as e: err = f"{type(e).__name__}: {e}"
        if not got:
            exts = sorted({f['filename'].rsplit('.', 1)[-1].lower() for f in fs})
            how = 'PENDING_PDF_DOC' if exts and not ({'xls','xlsx','xml'} & set(exts)) else 'PARSE_FAILED'
        status[rid] = dict(rows=got, method=how, err=err,
                           title=rec.get('Title',''), year=rec.get('YearTextMonthOrPeriod',''),
                           exts=sorted({f['filename'].rsplit('.',1)[-1].lower() for f in fs}))

    # dedupe: exact first, then logical (same key + same raw value)
    seen, exact = set(), 0
    uniq = []
    for r in rows:
        k = tuple(str(r[c]) for c in COLS)
        if k in seen: exact += 1; continue
        seen.add(k); uniq.append(r)
    lseen, logical = set(), 0
    final = []
    for r in uniq:
        k = tuple(str(r[c]) for c in KEY) + (str(r['Giá trị gốc']),)
        if k in lseen: logical += 1; continue
        lseen.add(k); final.append(r)
    return final, status, dict(exact_dup=exact, logical_dup=logical, raw=len(rows))

def write(province, final, status, stats, dest):
    wb = Workbook(); ws = wb.active; ws.title = 'Data'
    ws.append([f"{province.upper()} – DỮ LIỆU CÔNG KHAI NGÂN SÁCH (nguồn: CKNS Bộ Tài chính)"])
    ws.append([f"Sinh tự động {datetime.date.today():%d/%m/%Y} | {len(final)} dòng | "
               f"blank giữ blank, '-' giữ '-', không suy đoán"])
    ws.append(COLS)
    for c in ws[3]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in final: ws.append([r[c] for c in COLS])
    for i, w in enumerate([16,16,20,12,14,14,20,40,40,28,14,16,12,18,46,34], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = 'A4'

    q = wb.create_sheet('Summary_QA')
    meth = collections.Counter(v['method'] for v in status.values())
    num = sum(1 for r in final if r['Giá trị chuẩn hóa'] != '')
    vnd = sum(1 for r in final if r['Quy đổi VND'] != '')
    rowsum = [('Tỉnh/TP', province), ('Nguồn', 'CKNS Bộ Tài chính – DeptService.svc/SearchReport'),
              ('Số báo cáo (report)', len(status)),
              ('Báo cáo trích được dữ liệu', sum(1 for v in status.values() if v['rows'])),
              ('Báo cáo theo XML', meth.get('XML', 0)), ('Báo cáo theo XLS/XLSX', meth.get('XLSX', 0)),
              ('Báo cáo chờ PDF/DOC (pending)', meth.get('PENDING_PDF_DOC', 0)),
              ('Báo cáo parse lỗi', meth.get('PARSE_FAILED', 0)),
              ('Dòng thô', stats['raw']), ('Trùng tuyệt đối đã loại', stats['exact_dup']),
              ('Trùng logic đã loại', stats['logical_dup']), ('Dòng cuối cùng', len(final)),
              ('Dòng có giá trị chuẩn hóa', num), ('Dòng blank giữ nguyên', len(final) - num),
              ('Dòng có Quy đổi VND', vnd)]
    q.append([f'TỔNG HỢP NGUỒN & QA – {province.upper()}']); q.append([])
    q.append(['Chỉ tiêu', 'Giá trị'])
    for c in q[3]: c.font = Font(bold=True)
    for k, v in rowsum: q.append([k, v])
    q.column_dimensions['A'].width = 34; q.column_dimensions['B'].width = 56

    p = wb.create_sheet('Pending_Review')
    p.append(['report_id', 'năm', 'tiêu đề', 'định dạng có sẵn', 'trạng thái', 'lỗi'])
    for c in p[1]: c.font = Font(bold=True)
    for rid, v in sorted(status.items()):
        if v['method'] in ('PENDING_PDF_DOC', 'PARSE_FAILED'):
            p.append([rid, v['year'], v['title'], ','.join(v['exts']), v['method'], v['err'][:180]])
    for col, w in zip('ABCDEF', [11, 8, 60, 20, 18, 50]): p.column_dimensions[col].width = w
    wb.save(dest); return dest

if __name__ == '__main__':
    prov, wd, dest = sys.argv[1], sys.argv[2], sys.argv[3]
    final, status, stats = build(prov, wd)
    write(prov, final, status, stats, dest)
    print(f"rows={len(final)} raw={stats['raw']} dup_exact={stats['exact_dup']} dup_logic={stats['logical_dup']}")
    print("methods:", collections.Counter(v['method'] for v in status.values()).most_common())
    print("->", dest)
