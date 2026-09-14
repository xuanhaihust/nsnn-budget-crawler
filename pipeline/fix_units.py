"""Bring the ĐVT column of existing workbooks in line with the current unit rules.

    .venv/bin/python pipeline/fix_units.py --check           # report, write nothing
    .venv/bin/python pipeline/fix_units.py --apply           # rewrite output/*.xlsx
    .venv/bin/python pipeline/fix_units.py --apply "Phú Thọ" # just one province

No re-crawl, and no network. Everything this changes is a pure function of what the workbook
already contains:

    ĐVT            -> canon_unit(ĐVT)
    Quy đổi VND    -> Giá trị chuẩn hóa x unit_factor(ĐVT)
    Ghi chú        -> "; ĐVT nguồn: <as published>" appended when the spelling changed

So `./nsnn` is only needed when the SOURCE data changes. Re-downloading several GB to restate
a unit string would be work for nothing - the raw files under work/ have no say in any of the
three columns above.

Two rules this must not break, and it checks both before saving:

* A blank stays blank. If `Giá trị chuẩn hóa` is empty, `Quy đổi VND` stays empty, whatever
  the unit says.
* An existing conversion is never restated. If a row already has a VND figure and the new
  factor would give a different one, the file is left alone and the row is reported - that
  would be changing a published number, not tidying a spelling.
"""
import os, pathlib, sys
from concurrent.futures import ProcessPoolExecutor

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from openpyxl import load_workbook                              # noqa: E402
from parse_xml import canon_unit, unit_factor                   # noqa: E402
import harvest                                                  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / 'output'
COLS = ('ĐVT', 'Giá trị chuẩn hóa', 'Quy đổi VND', 'Ghi chú')


def header_row(ws):
    """Rows 1-2 are a title banner; the column names sit below them."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
        if row and 'ĐVT' in row:
            return i, {name: row.index(name) for name in COLS if name in row}
    return None, {}


def fix_sheet(ws, apply):
    """Returns (rows_seen, spelling_changes, conversions_added, conflicts)."""
    hrow, idx = header_row(ws)
    if hrow is None or 'ĐVT' not in idx:
        return 0, 0, 0, []
    u_i, s_i = idx['ĐVT'], idx.get('Giá trị chuẩn hóa')
    v_i, g_i = idx.get('Quy đổi VND'), idx.get('Ghi chú')
    seen = renamed = gained = 0
    conflicts = []
    for r in range(hrow + 1, ws.max_row + 1):
        src = ws.cell(r, u_i + 1).value
        src = src.strip() if isinstance(src, str) else ''
        seen += 1
        canon = canon_unit(src)
        factor = unit_factor(src)
        std = ws.cell(r, s_i + 1).value if s_i is not None else None
        old_vnd = ws.cell(r, v_i + 1).value if v_i is not None else None

        want = None
        if factor and isinstance(std, (int, float)):
            want = std * factor                       # blank stays blank: std must be a number
        # Relative, not exact: Excel stores 2110151000 where value*factor recomputes to
        # 2110150999.9999998, which is float noise, not a restatement. A real unit conflict is
        # 1000x or 1e6x apart, so 1e-9 still catches every one of them.
        if isinstance(old_vnd, (int, float)) and want is not None and \
                abs(old_vnd - want) > max(abs(want), abs(old_vnd)) * 1e-9:
            conflicts.append((r, src, old_vnd, want))
            continue
        if canon != src:
            renamed += 1
            if apply:
                ws.cell(r, u_i + 1).value = canon
                if g_i is not None:
                    note = ws.cell(r, g_i + 1).value or ''
                    if 'ĐVT nguồn:' not in str(note):
                        ws.cell(r, g_i + 1).value = f"{note}; ĐVT nguồn: {src}" if note \
                            else f"ĐVT nguồn: {src}"
        if want is not None and not isinstance(old_vnd, (int, float)):
            gained += 1
            if apply and v_i is not None:
                ws.cell(r, v_i + 1).value = want
    return seen, renamed, gained, conflicts


def _one(job):
    """Process one workbook in its own process. Returns a picklable summary."""
    f, apply = job
    wb = load_workbook(f)
    rows, ren, gain, conflicts = fix_sheet(wb['Data'], apply)
    if apply and not conflicts and (ren or gain):
        wb.save(f)
    wb.close()
    return f.name, rows, ren, gain, [(f.name, *c) for c in conflicts]


def main():
    apply = '--apply' in sys.argv
    if not apply and '--check' not in sys.argv:
        sys.exit(__doc__)
    wanted = [a for a in sys.argv[1:] if not a.startswith('--')]
    files = sorted(OUT.glob('*_budget_CKNS_*.xlsx'))
    if wanted:
        slugs = {harvest.slug(w) for w in wanted}
        files = [f for f in files if any(f.name.startswith(s + '_budget') for s in slugs)]
    if not files:
        sys.exit('no matching workbooks in output/')

    tot_rows = tot_ren = tot_gain = 0
    all_conflicts = []
    # One process per workbook, because openpyxl spends nearly all of its time parsing and
    # re-serialising XML and that is CPU-bound. Serially this is a ~5 hour job on the full set.
    workers = max(1, min(len(files), (os.cpu_count() or 2)))
    done = 0
    with ProcessPoolExecutor(max_workers=workers) as ex:
        for name, rows, ren, gain, conflicts in ex.map(_one, [(f, apply) for f in files]):
            tot_rows += rows; tot_ren += ren; tot_gain += gain
            all_conflicts += conflicts
            done += 1
            print(f"  [{done}/{len(files)}] {name:<46} {rows:>8,} rows  "
                  f"{ren:>7,} respelt  {gain:>5,} converted"
                  + ('  !! CONFLICT, not saved' if conflicts else ''), flush=True)

    print(f"\n{tot_rows:,} rows across {len(files)} workbooks")
    print(f"{tot_ren:,} unit spellings unified")
    print(f"{tot_gain:,} rows gained a Quy đổi VND")
    if all_conflicts:
        print(f"\n!! {len(all_conflicts)} rows already had a DIFFERENT VND figure. Those "
              "workbooks were NOT saved - restating a published conversion is not a cleanup:")
        for c in all_conflicts[:10]:
            print(f"   {c[0]} row {c[1]}: {c[2]!r} had {c[3]}, rule says {c[4]}")
    elif not apply:
        print("\n(--check: nothing was written)")


if __name__ == '__main__':
    main()
