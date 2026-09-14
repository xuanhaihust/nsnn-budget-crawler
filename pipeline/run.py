"""Crawl provincial budget data from CKNS.

    ./nsnn                          all 34 provinces, 5 at a time
    ./nsnn "Bắc Ninh" "Hưng Yên"    just these
    ./nsnn --jobs 3 --all           override the batch size
    ./nsnn --with-pdf "Hà Nội"      also store PDF/DOC (never parsed; for a future OCR pass)
    ./nsnn --list                   every CKNS department name and id
    ./nsnn --status                 what is already built
    ./nsnn --units                  every ĐVT spelling in output/, and which convert

Provinces run in separate processes. Each writes only to its own work/<slug>/ and its own
output file, so batches never collide. Downloads inside one province use 16 threads.

By default only machine-readable attachments (xml/xls/xlsx) are fetched - PDF/DOC are
never parsed and were 98% of the bytes. Reports that publish nothing else still appear in
Pending_Review. --with-pdf restores the old behaviour.
"""
import sys, pathlib, datetime, unicodedata, difflib, time, traceback
from concurrent.futures import ProcessPoolExecutor, as_completed

sys.path.insert(0, str(pathlib.Path(__file__).parent))
import ckns, harvest, build

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT, WORK = ROOT / 'output', ROOT / 'work'
LIST = ROOT / 'pipeline' / 'provinces34.txt'
JOBS = 5

# the Source Master uses post-merger names; CKNS still uses some older ones
ALIAS = {'Huế': 'Thừa Thiên Huế', 'TP Hồ Chí Minh': 'Hồ Chí Minh'}


def fold(s):
    s = unicodedata.normalize('NFD', str(s)).lower()
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def resolve(name, deps):
    name = ALIAS.get(name.strip(), name.strip())
    if name in deps:
        return name, deps[name]
    hit = [k for k in deps if fold(k) == fold(name)]
    if hit:
        return hit[0], deps[hit[0]]
    near = difflib.get_close_matches(fold(name), [fold(k) for k in deps], n=5, cutoff=0.6)
    raise LookupError(f"province {name!r} not found; closest: {near} (try --list)")


def provinces34():
    return [l.strip() for l in LIST.read_text().splitlines() if l.strip()]


def work_one(title, dept_id, with_pdf=False):
    """Runs in its own process. Returns a summary dict; never raises past the caller."""
    t0 = time.time()
    try:
        harvest.run(title, dept_id, WORK, with_pdf=with_pdf)
        OUT.mkdir(exist_ok=True)
        dest = OUT / f"{harvest.slug(title)}_budget_CKNS_{datetime.date.today():%Y-%m-%d}.xlsx"
        final, status, stats = build.build(title, WORK / harvest.slug(title))
        build.write(title, final, status, stats, dest)
        return dict(province=title, ok=True, rows=len(final),
                    reports=len(status), parsed=sum(1 for v in status.values() if v['rows']),
                    pending=sum(1 for v in status.values() if not v['rows']),
                    dest=str(dest.relative_to(ROOT)), secs=round(time.time() - t0))
    except Exception as e:
        return dict(province=title, ok=False, err=f"{type(e).__name__}: {e}",
                    tb=traceback.format_exc()[-400:], secs=round(time.time() - t0))


def run(names, jobs=JOBS, with_pdf=False):
    deps = ckns.departments()
    targets, unknown = [], []
    for n in names:
        try:
            targets.append(resolve(n, deps))
        except LookupError as e:
            unknown.append(str(e))
    for u in unknown:
        print(f"SKIP  {u}")
    if not targets:
        return 1

    print(f"Running {len(targets)} province(s), {jobs} at a time\n")
    done, results = 0, []
    t0 = time.time()
    with ProcessPoolExecutor(max_workers=jobs) as ex:
        futs = {ex.submit(work_one, t, d, with_pdf): t for t, d in targets}
        for f in as_completed(futs):
            r = f.result(); results.append(r); done += 1
            tag = f"[{done}/{len(targets)}]"
            if r['ok']:
                print(f"{tag} OK    {r['province']:<18} {r['rows']:>7} rows | "
                      f"{r['parsed']}/{r['reports']} reports | {r['secs']}s -> {r['dest']}")
            else:
                print(f"{tag} FAIL  {r['province']:<18} {r['err']}")

    good = [r for r in results if r['ok']]
    bad = [r for r in results if not r['ok']]
    print(f"\n{'='*64}\n{len(good)} ok, {len(bad)} failed, {len(unknown)} unknown | "
          f"{sum(r['rows'] for r in good):,} rows | {round(time.time()-t0)}s total")
    for r in bad:
        print(f"  FAILED {r['province']}: {r['err']}")
    for r in good:
        if r['pending']:
            print(f"  note   {r['province']}: {r['pending']} report(s) in Pending_Review")
    return 1 if bad else 0


def status():
    built = {p.name.split('_budget_CKNS_')[0]: p for p in sorted(OUT.glob('*_budget_CKNS_*.xlsx'))}
    names = provinces34()
    print(f"{len(built)}/{len(names)} provinces built\n")
    for n in names:
        s = harvest.slug(ALIAS.get(n, n))
        p = built.get(s)
        print(f"  {'DONE' if p else '    '}  {n:<20} {p.name if p else ''}")
    missing = [n for n in names if harvest.slug(ALIAS.get(n, n)) not in built]
    if missing:
        print(f"\nremaining {len(missing)}: " + ', '.join(missing))


def units():
    """Every distinct ĐVT spelling across output/, with the factor each resolves to.

    A standing check, not a one-off cleanup. Unit spellings are the most reliably recurring
    defect in this corpus - a misspelling silently costs those rows their Quy đổi VND, and it
    has happened three separate times. Anything listed as NOT CONVERTED that looks like money
    belongs in parse_xml.CURRENCY or SCALE; everything else on that list is correctly left
    alone (percentages, counts, and the odd agency name).

    Reads the workbooks rather than the warehouse, because the workbook is what is handed over.
    """
    import collections
    from openpyxl import load_workbook
    from parse_xml import unit_factor, canon_unit, CURRENCY

    seen, files = collections.Counter(), sorted(OUT.glob('*_budget_CKNS_*.xlsx'))
    if not files:
        print('no workbooks in output/ - run ./nsnn first')
        return
    for n, f in enumerate(files, 1):
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb['Data']
        col = first = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
            if row and 'ĐVT' in row:      # rows 1-2 are a title banner, not the header
                col, first = row.index('ĐVT'), i + 1
                break
        if col is None:
            print(f"  !! {f.name}: no ĐVT column, skipped")
            wb.close()
            continue
        for row in ws.iter_rows(min_row=first, values_only=True):
            v = row[col] if col < len(row) else None
            seen[v.strip() if isinstance(v, str) else ''] += 1
        wb.close()
        print(f"  [{n}/{len(files)}] {f.name}", flush=True)

    print(f"\n{len(seen)} distinct ĐVT spellings across {len(files)} workbooks\n")
    conv = [(u, c) for u, c in seen.items() if unit_factor(u) is not None]
    rest = [(u, c) for u, c in seen.items() if unit_factor(u) is None]
    print(f"{'ĐVT':<24} {'factor':>14} {'rows':>10}  canonical")
    print('-' * 68)
    for u, c in sorted(conv, key=lambda x: -x[1]):
        flag = '' if canon_unit(u) == u else f"  <- written as {canon_unit(u)!r}"
        print(f"{u!r:<24} {unit_factor(u):>14,} {c:>10,}{flag}")
    print(f"\nNOT CONVERTED (Quy đổi VND stays blank, by design):")
    for u, c in sorted(rest, key=lambda x: -x[1]):
        print(f"  {u!r:<24} {c:>10,}")
    bad = [u for u, _ in rest if u and any(w in u.lower() for w in CURRENCY)]
    if bad:
        print(f"\n!! {len(bad)} of those name đồng but do not convert - add the spelling to "
              f"parse_xml.CURRENCY/SCALE: {bad}")


def main(argv):
    args, jobs, names = argv[:], JOBS, []
    with_pdf = '--with-pdf' in args
    if with_pdf:
        args.remove('--with-pdf')
    if '--jobs' in args:
        i = args.index('--jobs'); jobs = int(args[i + 1]); del args[i:i + 2]
    if '--list' in args:
        for k, v in sorted(ckns.departments().items()):
            print(f"{v:>4}  {k}")
        return 0
    if '--status' in args:
        status(); return 0
    if '--units' in args:
        units(); return 0
    names = [a for a in args if a != '--all']
    if not names:                       # no province given means every province
        names = provinces34()
    return run(names, jobs, with_pdf)


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
